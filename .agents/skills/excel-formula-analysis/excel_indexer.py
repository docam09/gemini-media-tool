#!/usr/bin/env python3
"""Excel formula & dependency indexer.

Quickly extracts every formula, cached value, and cross-sheet relationship from
a multi-sheet Excel workbook (.xlsx / .xlsm / .xlsb / .xls) so an AI agent can
explain the meaning of any cell on demand.

Workflow
--------
1. `build`  : parse the workbook once and write a compact `index.json` plus a
              human-readable `summary.md` next to the workbook (or to --out).
2. `cell`   : explain a single cell (formula, value, precedents, dependents).
3. `find`   : search formulas / comments / values for a substring.
4. `sheet`  : summarise one sheet (size, formula count, inbound/outbound links).

Why a build step? Re-parsing a large workbook for every question is slow. The
`build` step does the heavy parsing once; the other commands read `index.json`,
so answering "what does cell X mean?" is instant.

Reading engines (pick with `build --engine`)
--------------------------------------------
openpyxl cannot read the binary .xlsb / .xls formats, and pyxlsb only exposes
cached *values* (not formula strings). So `build` supports several engines:

* `openpyxl`    : reads .xlsx / .xlsm directly (no Excel needed).
* `com`         : Windows + Microsoft Excel via pywin32. Opens the workbook
                  (including .xlsb) read-only and reads `.Formula` / `.Value`
                  DIRECTLY -- NO conversion, NO temp file written to disk.
                  Best when you have Excel and cannot write converted files
                  (e.g. locked-down C: drive).
* `libreoffice` : converts the binary workbook to .xlsx via LibreOffice headless
                  (`soffice --headless --convert-to xlsx`) then reads it. Use
                  `--workdir` to control where the temp .xlsx is written (e.g. a
                  drive you can write to). LibreOffice recalculates and stores
                  cached values during conversion, so both formulas and values
                  survive.

`--engine auto` (default) chooses: openpyxl for .xlsx/.xlsm; otherwise Excel COM
if available on Windows, else LibreOffice.
"""

from __future__ import annotations

import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile
from collections import defaultdict

from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer, Token
from openpyxl.utils import range_boundaries, get_column_letter
from openpyxl.utils.cell import coordinate_from_string

# A referenced range larger than this many cells is recorded as a "bulk"
# reference instead of being expanded cell-by-cell (avoids blow-ups on
# whole-column refs like A:A or huge table ranges).
MAX_RANGE_EXPANSION = 4096


# --------------------------------------------------------------------------- #
# Loading / conversion
# --------------------------------------------------------------------------- #
def find_soffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        path = shutil.which(name)
        if path:
            return path
    return None


def convert_to_xlsx(path: str, workdir: str | None = None) -> str:
    """Convert a binary workbook (.xlsb/.xls) to .xlsx via LibreOffice headless.

    ``workdir`` controls where the converted .xlsx (and LibreOffice profile) are
    written -- point it at a writable location if the default temp dir / C:
    drive is locked down. Returns the path to the produced .xlsx file.
    """
    soffice = find_soffice()
    if not soffice:
        sys.exit(
            "ERROR: this file needs LibreOffice to read formulas, but neither "
            "'soffice' nor 'libreoffice' is on PATH.\n"
            "Install it (e.g. `sudo apt-get install -y libreoffice-calc`), or on "
            "Windows with Excel use `--engine com` instead (no conversion)."
        )
    if workdir:
        os.makedirs(workdir, exist_ok=True)
    out_dir = tempfile.mkdtemp(prefix="xlsx_convert_", dir=workdir)
    # A dedicated profile dir avoids clashes with any running LibreOffice.
    profile = tempfile.mkdtemp(prefix="lo_profile_", dir=workdir)
    profile_url = "file:///" + os.path.abspath(profile).replace(os.sep, "/").lstrip("/")
    cmd = [
        soffice,
        "-env:UserInstallation=" + profile_url,
        "--headless",
        "--calc",
        "--convert-to",
        "xlsx",
        "--outdir",
        out_dir,
        path,
    ]
    proc = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
    base = os.path.splitext(os.path.basename(path))[0] + ".xlsx"
    produced = os.path.join(out_dir, base)
    if not os.path.exists(produced):
        sys.exit(
            "ERROR: LibreOffice conversion failed.\n"
            f"stdout: {proc.stdout}\nstderr: {proc.stderr}"
        )
    return produced


def _com_available() -> bool:
    if not sys.platform.startswith("win"):
        return False
    try:
        import win32com.client  # noqa: F401
    except ImportError:
        return False
    return True


def choose_engine(path: str, requested: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext not in (".xlsx", ".xlsm", ".xlsb", ".xls"):
        sys.exit(f"ERROR: unsupported extension '{ext}'. Use xlsx/xlsm/xlsb/xls.")
    if requested != "auto":
        return requested
    if ext in (".xlsx", ".xlsm"):
        return "openpyxl"
    if _com_available():
        return "com"
    if find_soffice():
        return "libreoffice"
    return "com" if sys.platform.startswith("win") else "libreoffice"


# --------------------------------------------------------------------------- #
# Formula reference parsing
# --------------------------------------------------------------------------- #
def split_sheet_ref(ref: str) -> tuple[str | None, str]:
    """Split a range token like ``'My Sheet'!$A$1:$B$2`` into (sheet, range).

    Returns (None, range) when there is no sheet qualifier.
    """
    if "!" not in ref:
        return None, ref
    sheet_part, _, rng = ref.rpartition("!")
    sheet = sheet_part.strip()
    if sheet.startswith("'") and sheet.endswith("'"):
        sheet = sheet[1:-1].replace("''", "'")
    # External-workbook refs look like [1]Sheet or [book.xlsx]Sheet -> keep raw.
    return sheet, rng


def expand_range(sheet: str, rng: str) -> tuple[list[str], bool]:
    """Expand ``rng`` on ``sheet`` to a list of ``Sheet!A1`` coords.

    Returns (coords, is_bulk). For whole-column/row refs or ranges larger than
    MAX_RANGE_EXPANSION, returns ([], True) so callers treat it as a bulk ref.
    """
    rng = rng.replace("$", "")
    try:
        min_col, min_row, max_col, max_row = range_boundaries(rng)
    except Exception:
        return [], False
    # Whole-column (A:A) or whole-row (1:1) refs have None boundaries.
    if None in (min_col, min_row, max_col, max_row):
        return [], True
    n = (max_col - min_col + 1) * (max_row - min_row + 1)
    if n > MAX_RANGE_EXPANSION:
        return [], True
    coords = []
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            coords.append(f"{sheet}!{get_column_letter(c)}{r}")
    return coords, False


def extract_references(formula: str, current_sheet: str) -> tuple[set[str], list[str]]:
    """Parse a formula and return (precedent_cells, bulk_refs).

    precedent_cells : set of fully-qualified ``Sheet!A1`` coordinates.
    bulk_refs       : list of fully-qualified range strings too large to expand.
    """
    cells: set[str] = set()
    bulk: list[str] = []
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return cells, bulk
    for tok in tokens:
        if tok.type != Token.OPERAND or tok.subtype != Token.RANGE:
            continue
        sheet, rng = split_sheet_ref(tok.value)
        target_sheet = sheet if sheet is not None else current_sheet
        coords, is_bulk = expand_range(target_sheet, rng)
        if is_bulk:
            bulk.append(f"{target_sheet}!{rng.replace('$', '')}")
        elif coords:
            cells.update(coords)
        else:
            # Could be a defined-name token; record verbatim as a precedent.
            bulk.append(tok.value)
    return cells, bulk


def _jsonable(v):
    """Coerce a cell value into something json-serialisable."""
    import datetime

    if v is None or isinstance(v, (str, int, float, bool)):
        return v
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        return v.isoformat()
    try:
        json.dumps(v)
        return v
    except (TypeError, ValueError):
        return str(v)


# --------------------------------------------------------------------------- #
# Extraction engines
# --------------------------------------------------------------------------- #
# Each engine returns a "raw" tuple that the engine-independent `analyze` step
# turns into the final index:
#   (sheet_order, sheets_meta, named_ranges, cells_raw)
# where cells_raw maps "Sheet!A1" -> {"formula"?, "value"?, "comment"?}.

def extract_openpyxl(xlsx_path: str) -> tuple[list, dict, dict, dict]:
    wb_f = load_workbook(xlsx_path, data_only=False, read_only=False)
    wb_v = load_workbook(xlsx_path, data_only=True, read_only=False)

    cells_raw: dict[str, dict] = {}
    sheets_meta: dict[str, dict] = {}

    for ws in wb_f.worksheets:
        name = ws.title
        wv = wb_v[name]
        formula_count = 0
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if val is None and cell.comment is None:
                    continue
                coord = f"{name}!{cell.coordinate}"
                entry: dict = {}
                if isinstance(val, str) and val.startswith("="):
                    formula_count += 1
                    entry["formula"] = val
                    cached = wv[cell.coordinate].value
                    if cached is not None:
                        entry["value"] = _jsonable(cached)
                elif val is not None:
                    entry["value"] = _jsonable(val)
                if cell.comment is not None:
                    entry["comment"] = cell.comment.text
                if entry:
                    cells_raw[coord] = entry
        sheets_meta[name] = {
            "max_row": ws.max_row or 0,
            "max_col": ws.max_column or 0,
            "dimensions": ws.dimensions,
            "formula_count": formula_count,
            "state": ws.sheet_state,  # visible / hidden / veryHidden
        }

    named: dict[str, str] = {}
    try:
        for name_obj in wb_f.defined_names.values():
            named[name_obj.name] = name_obj.value
    except AttributeError:
        for key in wb_f.defined_names:
            try:
                named[key] = wb_f.defined_names[key].value
            except Exception:
                pass

    return wb_f.sheetnames, sheets_meta, named, cells_raw


def extract_com(path: str) -> tuple[list, dict, dict, dict]:
    """Read formulas/values directly via Microsoft Excel (pywin32) -- no temp
    file written to disk. Windows + Excel only."""
    try:
        import pythoncom
        import win32com.client as win32
    except ImportError:
        sys.exit(
            "ERROR: the 'com' engine requires Microsoft Excel and pywin32 on "
            "Windows.\n  pip install pywin32"
        )

    state_map = {-1: "visible", 0: "hidden", 2: "veryHidden"}
    cells_raw: dict[str, dict] = {}
    sheets_meta: dict[str, dict] = {}
    sheet_order: list[str] = []
    named: dict[str, str] = {}

    pythoncom.CoInitialize()
    # DispatchEx starts a private Excel instance so Quit() never closes a
    # workbook the user already had open.
    excel = win32.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        # Force A1 style so .Formula is parseable even if Excel is set to R1C1.
        excel.ReferenceStyle = 1  # xlA1
    except Exception:
        pass
    try:
        wb = excel.Workbooks.Open(
            os.path.abspath(path), ReadOnly=True, UpdateLinks=0
        )
        try:
            for ws in wb.Worksheets:
                name = ws.Name
                sheet_order.append(name)
                used = ws.UsedRange
                first_row, first_col = used.Row, used.Column
                n_rows, n_cols = used.Rows.Count, used.Columns.Count
                formulas = used.Formula
                values = used.Value
                # A 1x1 UsedRange returns scalars rather than nested tuples.
                if n_rows == 1 and n_cols == 1:
                    formulas = ((formulas,),)
                    values = ((values,),)
                formula_count = 0
                for i in range(n_rows):
                    frow, vrow = formulas[i], values[i]
                    for j in range(n_cols):
                        f, v = frow[j], vrow[j]
                        if f is None and v is None:
                            continue
                        coord = (
                            f"{name}!{get_column_letter(first_col + j)}"
                            f"{first_row + i}"
                        )
                        entry: dict = {}
                        if isinstance(f, str) and f.startswith("="):
                            entry["formula"] = f
                            formula_count += 1
                            if v is not None:
                                entry["value"] = _jsonable(v)
                        elif v is not None:
                            entry["value"] = _jsonable(v)
                        elif f is not None:
                            entry["value"] = _jsonable(f)
                        if entry:
                            cells_raw[coord] = entry
                # Classic (non-threaded) comments.
                try:
                    for cm in ws.Comments:
                        addr = cm.Parent.Address(False, False)
                        ckey = f"{name}!{addr}"
                        cells_raw.setdefault(ckey, {})["comment"] = cm.Text()
                except Exception:
                    pass
                try:
                    state = state_map.get(int(ws.Visible), "visible")
                except Exception:
                    state = "visible"
                sheets_meta[name] = {
                    "max_row": first_row + n_rows - 1,
                    "max_col": first_col + n_cols - 1,
                    "dimensions": used.Address(False, False),
                    "formula_count": formula_count,
                    "state": state,
                }
            try:
                for nm in wb.Names:
                    try:
                        named[str(nm.Name)] = str(nm.RefersTo)
                    except Exception:
                        pass
            except Exception:
                pass
        finally:
            wb.Close(SaveChanges=False)
    finally:
        excel.Quit()
        pythoncom.CoUninitialize()

    return sheet_order, sheets_meta, named, cells_raw


# --------------------------------------------------------------------------- #
# Build (engine dispatch + engine-independent analysis)
# --------------------------------------------------------------------------- #
def analyze(sheet_order, sheets_meta, named, cells_raw, source, engine, converted):
    """Add precedents / dependents / sheet links to the raw cell data."""
    sheetnames = set(sheet_order)
    cells = cells_raw
    sheet_links: dict[str, set] = defaultdict(set)

    for coord, entry in cells.items():
        formula = entry.get("formula")
        if not formula:
            continue
        name = coord.rsplit("!", 1)[0]
        precedents, bulk = extract_references(formula, name)
        if precedents:
            entry["precedents"] = sorted(precedents)
        if bulk:
            entry["precedent_ranges"] = bulk
        for p in precedents:
            ps = p.rsplit("!", 1)[0]
            if ps != name:
                sheet_links[name].add(ps)
        for b in bulk:
            if "!" in b:
                bs = b.rsplit("!", 1)[0].strip("'")
                if bs != name and bs in sheetnames:
                    sheet_links[name].add(bs)

    # Reverse dependency map (dependents).
    dependents: dict[str, set] = defaultdict(set)
    for coord, entry in cells.items():
        for p in entry.get("precedents", []):
            dependents[p].add(coord)
    for coord, deps in dependents.items():
        if coord in cells:
            cells[coord]["dependents"] = sorted(deps)

    return {
        "source": os.path.abspath(source),
        "engine": engine,
        "converted_to_xlsx": converted,
        "sheet_order": sheet_order,
        "sheets": sheets_meta,
        "sheet_links": {k: sorted(v) for k, v in sheet_links.items()},
        "named_ranges": named,
        "cells": cells,
    }


def build_index(path: str, engine: str = "auto", workdir: str | None = None) -> dict:
    engine = choose_engine(path, engine)
    converted = False
    if engine == "openpyxl":
        sheet_order, sheets_meta, named, cells_raw = extract_openpyxl(path)
    elif engine == "com":
        sheet_order, sheets_meta, named, cells_raw = extract_com(path)
    elif engine == "libreoffice":
        xlsx_path = convert_to_xlsx(path, workdir)
        converted = True
        sheet_order, sheets_meta, named, cells_raw = extract_openpyxl(xlsx_path)
    else:
        sys.exit(f"ERROR: unknown engine '{engine}'.")
    return analyze(
        sheet_order, sheets_meta, named, cells_raw, path, engine, converted
    )


# --------------------------------------------------------------------------- #
# Reporting
# --------------------------------------------------------------------------- #
def write_summary(index: dict, out_md: str) -> None:
    lines = []
    lines.append(f"# Workbook index: {os.path.basename(index['source'])}")
    lines.append("")
    lines.append(f"- Source: `{index['source']}`")
    lines.append(f"- Engine: {index.get('engine', 'openpyxl')}")
    if index["converted_to_xlsx"]:
        lines.append("- Converted from binary format via LibreOffice.")
    total_formulas = sum(s["formula_count"] for s in index["sheets"].values())
    lines.append(f"- Sheets: {len(index['sheets'])}")
    lines.append(f"- Total formula cells: {total_formulas}")
    lines.append(f"- Named ranges: {len(index['named_ranges'])}")
    lines.append("")

    lines.append("## Sheets")
    lines.append("")
    lines.append("| # | Sheet | State | Used range | Formula cells | Reads from |")
    lines.append("|---|-------|-------|-----------|--------------|------------|")
    for i, name in enumerate(index["sheet_order"], 1):
        m = index["sheets"][name]
        reads = ", ".join(index["sheet_links"].get(name, [])) or "—"
        lines.append(
            f"| {i} | {name} | {m['state']} | {m['dimensions']} "
            f"| {m['formula_count']} | {reads} |"
        )
    lines.append("")

    # Sheet dependency graph (who feeds whom).
    lines.append("## Sheet dependency graph")
    lines.append("")
    lines.append("Arrows mean *reads data from* (A -> B: A's formulas reference B).")
    lines.append("")
    any_link = False
    for name in index["sheet_order"]:
        targets = index["sheet_links"].get(name, [])
        if targets:
            any_link = True
            lines.append(f"- `{name}` -> " + ", ".join(f"`{t}`" for t in targets))
    if not any_link:
        lines.append("_No cross-sheet references detected._")
    lines.append("")

    if index["named_ranges"]:
        lines.append("## Named ranges")
        lines.append("")
        for n, v in sorted(index["named_ranges"].items()):
            lines.append(f"- `{n}` = `{v}`")
        lines.append("")

    with open(out_md, "w") as f:
        f.write("\n".join(lines))


# --------------------------------------------------------------------------- #
# Query commands
# --------------------------------------------------------------------------- #
def load_index(index_path: str) -> dict:
    if not os.path.exists(index_path):
        sys.exit(
            f"ERROR: index '{index_path}' not found. Run the 'build' command first."
        )
    with open(index_path) as f:
        return json.load(f)


def cmd_cell(index: dict, coord: str, depth: int) -> None:
    coord = _canonical_coord(index, coord)
    cells = index["cells"]
    entry = cells.get(coord)
    print(f"# {coord}")
    if entry is None:
        print("(empty / no value or formula)")
        return
    if "formula" in entry:
        print(f"Formula : {entry['formula']}")
    if "value" in entry:
        print(f"Value   : {entry['value']!r}")
    if "comment" in entry:
        print(f"Comment : {entry['comment']}")
    if entry.get("precedents"):
        print("\nDepends on (precedents):")
        for p in entry["precedents"]:
            pv = cells.get(p, {})
            extra = pv.get("formula") or repr(pv.get("value"))
            print(f"  - {p}  =>  {extra}")
    if entry.get("precedent_ranges"):
        print("\nDepends on (ranges / names):")
        for b in entry["precedent_ranges"]:
            print(f"  - {b}")
    if entry.get("dependents"):
        print("\nUsed by (dependents):")
        for d in entry["dependents"]:
            dv = cells.get(d, {})
            print(f"  - {d}  =>  {dv.get('formula', '')}")
    if depth > 1:
        print("\n--- Precedent trace ---")
        _trace(cells, coord, depth, seen=set(), indent=0)


def _trace(cells, coord, depth, seen, indent):
    if coord in seen or depth <= 0:
        return
    seen.add(coord)
    entry = cells.get(coord, {})
    label = entry.get("formula") or repr(entry.get("value"))
    print("  " * indent + f"{coord}: {label}")
    for p in entry.get("precedents", []):
        _trace(cells, p, depth - 1, seen, indent + 1)


def cmd_find(index: dict, needle: str, in_values: bool) -> None:
    needle_l = needle.lower()
    hits = 0
    for coord, entry in index["cells"].items():
        haystacks = [entry.get("formula", ""), entry.get("comment", "")]
        if in_values:
            haystacks.append(str(entry.get("value", "")))
        if any(needle_l in h.lower() for h in haystacks):
            shown = entry.get("formula") or repr(entry.get("value"))
            print(f"{coord}: {shown}")
            hits += 1
    print(f"\n{hits} match(es).")


def cmd_sheet(index: dict, name: str) -> None:
    if name not in index["sheets"]:
        sys.exit(f"ERROR: sheet '{name}' not found. Available: {index['sheet_order']}")
    m = index["sheets"][name]
    print(f"# Sheet: {name}")
    print(f"State        : {m['state']}")
    print(f"Used range   : {m['dimensions']} (rows={m['max_row']}, cols={m['max_col']})")
    print(f"Formula cells: {m['formula_count']}")
    print(f"Reads from   : {', '.join(index['sheet_links'].get(name, [])) or '—'}")
    feeds = [s for s, t in index["sheet_links"].items() if name in t]
    print(f"Feeds into   : {', '.join(feeds) or '—'}")


def _canonical_coord(index: dict, coord: str) -> str:
    """Normalise a user-supplied coord: add sheet if missing, upper-case col,
    strip $ signs."""
    coord = coord.replace("$", "")
    if "!" not in coord:
        # Default to the first sheet.
        coord = f"{index['sheet_order'][0]}!{coord}"
    sheet, ref = coord.rsplit("!", 1)
    sheet = sheet.strip().strip("'")
    try:
        col, row = coordinate_from_string(ref)
        ref = f"{col.upper()}{row}"
    except Exception:
        pass
    return f"{sheet}!{ref}"


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #
def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="command", required=True)

    b = sub.add_parser("build", help="Parse workbook -> index.json + summary.md")
    b.add_argument("file")
    b.add_argument("--out", help="Output dir (default: alongside the workbook)")
    b.add_argument(
        "--engine",
        choices=["auto", "openpyxl", "com", "libreoffice"],
        default="auto",
        help="Reading engine. 'com' = MS Excel via pywin32 (no conversion, "
        "Windows only); 'libreoffice' = convert to xlsx first.",
    )
    b.add_argument(
        "--workdir",
        help="Writable dir for the LibreOffice conversion output (use a drive "
        "you can write to if C:/temp is locked down).",
    )

    c = sub.add_parser("cell", help="Explain a single cell")
    c.add_argument("--index", default="index.json")
    c.add_argument("coord", help="e.g. 'Summary!B2' or 'B2'")
    c.add_argument("--depth", type=int, default=1, help="Precedent trace depth")

    f = sub.add_parser("find", help="Search formulas/comments (and optionally values)")
    f.add_argument("--index", default="index.json")
    f.add_argument("needle")
    f.add_argument("--values", action="store_true", help="also search cached values")

    s = sub.add_parser("sheet", help="Summarise one sheet")
    s.add_argument("--index", default="index.json")
    s.add_argument("name")

    args = ap.parse_args(argv)

    if args.command == "build":
        index = build_index(args.file, engine=args.engine, workdir=args.workdir)
        out_dir = args.out or os.path.dirname(os.path.abspath(args.file))
        os.makedirs(out_dir, exist_ok=True)
        idx_path = os.path.join(out_dir, "index.json")
        md_path = os.path.join(out_dir, "summary.md")
        with open(idx_path, "w", encoding="utf-8") as fh:
            json.dump(index, fh, ensure_ascii=False, indent=0)
        write_summary(index, md_path)
        n_cells = len(index["cells"])
        n_formulas = sum(s["formula_count"] for s in index["sheets"].values())
        print(f"Indexed {len(index['sheets'])} sheets, {n_cells} non-empty cells, "
              f"{n_formulas} formulas (engine: {index['engine']}).")
        print(f"  index   -> {idx_path}")
        print(f"  summary -> {md_path}")
        return

    index = load_index(args.index)
    if args.command == "cell":
        cmd_cell(index, args.coord, args.depth)
    elif args.command == "find":
        cmd_find(index, args.needle, args.values)
    elif args.command == "sheet":
        cmd_sheet(index, args.name)


if __name__ == "__main__":
    main()
